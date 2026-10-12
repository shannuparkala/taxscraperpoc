"""Current-year additions: spreadsheet in, AI life, engine figures, template workbook out.

The model is faked and only ever picks a class. Every amount asserted here
comes from the engine's rule tables. Invented assets only - the owner's
template (a real client's list) stays out of the repository; its layout is
what ``addition_report`` reproduces.

The report claims what a return claims, in the Code's order, per
jurisdiction (owner's decision 2026-10-11): Section 179 on the elected
assets, then bonus on the basis left, then MACRS on the rest - federal, and
the selected state's regime over the same elections.
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal as D
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import get_settings
from app.core.errors import LlmError
from app.services.depreciation import addition_import, addition_report
from app.services.depreciation.addition_report import AdditionAsset, build, to_xlsx
from app.services.depreciation.engine import InputError, UnsupportedYearError
from app.services.depreciation.life_predictor import LifePredictor, class_for_life
from app.services.depreciation.rules import get_rules
from app.services.depreciation.state import evaluate_state
from app.services.gemini_service import LlmResponse


@pytest.fixture()
def client():
    from app.main import app

    with TestClient(app) as test_client:
        yield test_client


def asset(asset_id: str, description: str, acquired: date, cost: str, asset_class: str,
          years: str | None = None, **kw: Any) -> AdditionAsset:
    return AdditionAsset(asset_id, description, acquired, D(cost), asset_class, years, **kw)


SAMPLE = [
    asset("101", "Shop Equipment - Two-Post Lift", date(2025, 4, 30), "12000.50", "other", "5"),
    asset("102", "Office Equipment - Copier", date(2025, 1, 9), "3000.00", "computers"),
    asset("103", "Furniture - Conference Table", date(2025, 2, 28), "7000.00", "furniture"),
    asset("104", "Qualified Improvement Property - Interior Remodel", date(2025, 6, 23), "100000.00", "qip"),
    asset("105", "Building Improvement - Structural Remodel (39-Year)", date(2025, 6, 23), "200000.00", "nonresidential"),
    asset("106", "Land Improvement - Parking Lot", date(2025, 6, 23), "50000.00", "land_improvements"),
]
#: The same list with neither election: plain MACRS on every asset.
NO_ELECTIONS = [replace(a, elect_179=False, elect_bonus=False) for a in SAMPLE]
#: Bonus only, no Section 179.
BONUS_ONLY = [replace(a, elect_179=False) for a in SAMPLE]


def by_id(report: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {r["asset_id"]: r for r in report["rows"]}


def fed(row: dict[str, Any]) -> tuple[str, str, str, str]:
    f = row["federal"]
    return f["sec179"], f["bonus"], f["macrs"], f["total"]


def workbook_bytes(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Worked examples: Sec 179 -> bonus on what is left -> MACRS on the rest
# ---------------------------------------------------------------------------
#
# Tax year 2026, every asset placed in service in Q1 (half-year convention,
# no mid-quarter). 2026 rules: Sec 179 limit $2,560,000, phase-out from
# $4,090,000, SUV cap $32,000 (Rev. Proc. 2025-32); bonus 100% for property
# acquired after Jan. 19, 2025 (Notice 2026-11); first-year MACRS 20% (5-year),
# 14.29% (7-year), 2.461% (39-year placed in service in January).

SUV = asset("1", "Heavy SUV (GVWR 6,500 lb)", date(2026, 3, 10), "90000", "heavy_vehicle")
MACHINE = asset("2", "CNC milling machine", date(2026, 3, 15), "100000", "machinery")
PRESS = asset("3", "Press line", date(2026, 3, 15), "4200000", "machinery")
WAREHOUSE = asset("4", "Warehouse building", date(2026, 1, 10), "390000", "nonresidential")


@pytest.mark.parametrize(
    ("assets", "expected"),
    [
        # 179 capped at the SUV limit; bonus 100% of the 58,000 left; nothing left for MACRS.
        ([SUV], ("32000.00", "58000.00", "0.00", "90000.00")),
        # 179 only: the 58,000 left goes to MACRS at 20% = 11,600.
        ([replace(SUV, elect_bonus=False)], ("32000.00", "0.00", "11600.00", "43600.00")),
        # Bonus only: 100% of cost.
        ([replace(MACHINE, elect_179=False)], ("0.00", "100000.00", "0.00", "100000.00")),
        # Neither: 7-year MACRS, 14.29%.
        ([replace(MACHINE, elect_179=False, elect_bonus=False)], ("0.00", "0.00", "14290.00", "14290.00")),
        # Phase-out: 2,560,000 - (4,200,000 - 4,090,000) = 2,450,000 of 179; bonus on the 1,750,000 left.
        ([PRESS], ("2450000.00", "1750000.00", "0.00", "4200000.00")),
        # Same, no bonus: MACRS on 1,750,000 at 14.29% = 250,075.
        ([replace(PRESS, elect_bonus=False)], ("2450000.00", "0.00", "250075.00", "2700075.00")),
        # A building: neither election is available; 39-year mid-month, January: 2.461%.
        ([WAREHOUSE], ("0.00", "0.00", "9597.90", "9597.90")),
    ],
)
def test_worked_examples_apply_179_then_bonus_then_macrs(assets: list[AdditionAsset],
                                                           expected: tuple[str, ...]) -> None:
    (row,) = build(assets, tax_year=2026, state=None, rules=get_rules())["rows"]
    assert fed(row) == expected


def test_a_bonus_rate_follows_the_acquisition_date_not_the_placed_in_service_year() -> None:
    """Acquired Jan. 9, 2025 (before the Jan. 20 line): the 40% phase-down rate."""
    rows = by_id(build(BONUS_ONLY, tax_year=2025, state=None, rules=get_rules()))
    assert rows["102"]["bonus_rate"] == "40"
    assert fed(rows["102"]) == ("0.00", "1200.00", "360.00", "1560.00")  # MACRS 20% of the 1,800 left
    assert rows["103"]["bonus_rate"] == "100" and rows["103"]["federal"]["bonus"] == "7000.00"
    assert rows["106"]["bonus_rate"] == "100"  # land improvements are qualified property


def test_macrs_uses_the_pub_946_first_year_rate_for_each_class() -> None:
    rows = by_id(build(NO_ELECTIONS, tax_year=2025, state=None, rules=get_rules()))
    assert rows["101"]["federal"]["macrs_rate"] == "20"
    assert rows["103"]["federal"]["macrs_rate"] == "14.29"
    assert rows["104"]["federal"]["macrs_rate"] == "3.33"  # QIP: 15-year straight line, not 150DB
    assert rows["106"]["federal"]["macrs_rate"] == "5"  # land improvements: 15-year 150DB
    assert rows["105"]["federal"]["macrs_rate"] == "1.391"  # 39-year, mid-month, placed in service in June
    assert rows["105"]["federal"]["macrs"] == "2782.00"
    # A cost with cents: the engine works on the basis in whole dollars, as a
    # return does, and the cell holds its amount rather than a formula.
    assert rows["101"]["federal"]["macrs"] == "2400.00"
    assert not rows["101"]["federal"]["macrs_formula"]


def test_totals_round_once_like_the_workbooks_sum_formulas() -> None:
    """Two 7-year assets at $1,001: 143.0429 each. Excel's SUM gives 286.09; summing rows
    already rounded to cents would show 286.08 on screen - a cent off the file."""
    pair = [asset(str(i), "Hoist", date(2025, 3, 1), "1001", "machinery", elect_179=False, elect_bonus=False)
            for i in (1, 2)]
    report = build(pair, tax_year=2025, state=None, rules=get_rules())
    assert [r["federal"]["macrs"] for r in report["rows"]] == ["143.04", "143.04"]
    assert report["totals"]["federal"]["macrs"] == "286.09"


def test_life_column_and_summary_by_life() -> None:
    report = build(NO_ELECTIONS, tax_year=2025, state=None, rules=get_rules())
    assert [r["life"] for r in report["rows"]] == ["5", "5", "7", "15", "39", "15"]
    summary = {s["life"]: s for s in report["summary"]}
    assert list(summary) == ["5", "7", "15", "39"]
    assert summary["15"]["cost"] == "150000.00"
    assert summary["15"]["depreciation"] == str(D("3330.00") + D("2500.00"))


def test_the_section_179_dollar_limit_and_phase_out_are_reported() -> None:
    report = build([PRESS], tax_year=2026, state=None, rules=get_rules())
    s179 = report["section_179"]
    assert (s179["limit"], s179["threshold"], s179["reduction"], s179["dollar_limit"]) == (
        "2560000", "4090000", "110000", "2450000")


# ---------------------------------------------------------------------------
# Elections and eligibility
# ---------------------------------------------------------------------------


def test_an_unticked_179_election_takes_no_section_179() -> None:
    (row,) = build([replace(MACHINE, elect_179=False, elect_bonus=False)], tax_year=2026, state=None,
                   rules=get_rules())["rows"]
    assert row["federal"]["sec179"] == "0.00" and row["elections"] == "MACRS only"


@pytest.mark.parametrize("state", [None, "CA", "MN"])
def test_report_bonus_selections_are_individual_and_mixed_classes_are_flagged(state: str | None) -> None:
    assets = [
        replace(MACHINE, elect_179=False),
        asset("5", "Forklift", date(2026, 3, 15), "30000", "machinery", elect_179=False, elect_bonus=False),
        asset("6", "Laptop", date(2026, 3, 15), "2000", "computers", elect_179=False),
    ]
    report = build(assets, tax_year=2026, state=state, rules=get_rules())
    rows = by_id(report)
    assert report["elect_out_classes"] == []
    warning = next(warning for warning in report["warnings"] if "mixed bonus selections for 7-year" in warning)
    assert "Planning only" in warning and "Before filing" in warning
    assert rows["2"]["elect_bonus"] and rows["2"]["bonus_applied"]
    assert fed(rows["2"]) == ("0.00", "100000.00", "0.00", "100000.00")
    assert not rows["5"]["bonus_applied"]
    assert fed(rows["5"]) == ("0.00", "0.00", "4287.00", "4287.00")
    assert rows["6"]["bonus_applied"] and rows["6"]["federal"]["bonus"] == "2000.00"
    assert report["totals"]["federal"]["total"] == "106287.00"
    for selected in assets:
        standalone = build([selected], tax_year=2026, state=state, rules=get_rules())["rows"][0]
        assert rows[selected.asset_id]["state"] == standalone["state"]
    worksheet = load_workbook(io.BytesIO(to_xlsx(report))).active
    assert warning in [cell.value for row in worksheet for cell in row]


def test_a_binding_179_limit_goes_first_where_bonus_would_not_cover_the_cost() -> None:
    """2025: the copier was acquired before Jan. 20 (40% bonus), the remodel after (100%).
    179 on the remodel adds nothing it would not get from bonus, so the limit covers the
    copier first - longest-lived-first alone left the copier at 40% bonus plus MACRS."""
    assets = [
        asset("1", "Qualified Improvement Property - Remodel", date(2025, 6, 23), "2500000.00", "qip"),
        asset("2", "Office Equipment - Copier", date(2025, 1, 9), "10000.00", "computers"),
    ]
    report = build(assets, tax_year=2025, state=None, rules=get_rules())
    rows = by_id(report)
    assert report["section_179"]["dollar_limit"] == "2500000"
    assert fed(rows["2"]) == ("10000.00", "0.00", "0.00", "10000.00")
    assert fed(rows["1"]) == ("2490000.00", "10000.00", "0.00", "2500000.00")
    assert report["totals"]["federal"]["total"] == "2510000.00"


def test_report_section_179_selections_are_individual() -> None:
    assets = [MACHINE, replace(MACHINE, asset_id="5", elect_179=False)]
    rows = build(assets, tax_year=2026, state=None, rules=get_rules())["rows"]
    assert fed(rows[0]) == ("100000.00", "0.00", "0.00", "100000.00")
    assert fed(rows[1]) == ("0.00", "100000.00", "0.00", "100000.00")


def test_ineligible_assets_say_why() -> None:
    rows = by_id(build(SAMPLE, tax_year=2025, state=None, rules=get_rules()))
    building, paving, copier = rows["105"], rows["106"], rows["102"]
    assert not building["bonus_eligible"] and "over 20 years" in building["bonus_reason"]
    assert not building["sec179_eligible"] and "not section 179 property" in building["sec179_reason"]
    assert paving["bonus_eligible"] and not paving["sec179_eligible"]
    assert "Land improvements" in paving["sec179_reason"]
    assert copier["sec179_eligible"] and copier["bonus_eligible"]
    assert copier["sec179_reason"] == copier["bonus_reason"] == ""
    # A ticked election an asset cannot take changes nothing.
    assert building["elect_179"] and building["federal"]["sec179"] == "0.00"


# ---------------------------------------------------------------------------
# The state group: the same elections under the state's regime
# ---------------------------------------------------------------------------


def test_california_allows_no_bonus_and_its_own_179_limit() -> None:
    """2026, CA verified: no 168(k); Section 179 limit $25,000."""
    bonus_only = build([replace(MACHINE, elect_179=False)], tax_year=2026, state="CA", rules=get_rules())
    (row,) = bonus_only["rows"]
    assert fed(row) == ("0.00", "100000.00", "0.00", "100000.00")
    assert row["state"] == {"sec179": "0", "bonus": "0", "macrs": "14290", "total": "14290", "addback": "0"}
    assert bonus_only["state"]["status"] == "computed"
    assert bonus_only["state"]["bonus_allowed"] is False
    assert "does not allow federal bonus" in bonus_only["state"]["bonus_note"]

    (both,) = build([MACHINE], tax_year=2026, state="CA", rules=get_rules())["rows"]
    assert fed(both) == ("100000.00", "0.00", "0.00", "100000.00")
    # 25,000 of state 179; 75,000 x 14.29% = 10,717.50, which the engine rounds to 10,718.
    assert both["state"] == {"sec179": "25000", "bonus": "0", "macrs": "10718", "total": "35718", "addback": "0"}


def test_delaware_allows_twenty_percent_for_2026_acquisitions() -> None:
    report = build([replace(MACHINE, elect_179=False)], tax_year=2026, state="DE", rules=get_rules())
    (row,) = report["rows"]
    assert row["state"] == {"sec179": "0", "bonus": "20000", "macrs": "11432", "total": "31432", "addback": "0"}
    assert report["state"]["bonus_allowed"] is True


def test_minnesota_adds_back_eighty_percent_of_federal_bonus() -> None:
    report = build([replace(MACHINE, elect_179=False)], tax_year=2026, state="MN", rules=get_rules())
    (row,) = report["rows"]
    assert row["state"]["addback"] == "80000" and row["state"]["total"] == "20000"
    assert "adds back 80%" in report["state"]["bonus_note"]


@pytest.mark.parametrize("code", ["CA", "WI", "MN"])
def test_the_state_total_is_the_states_own_first_year_deduction(code: str) -> None:
    """Per-asset state figures plus any year-level adjustment equal what the state engine says."""
    from app.services.depreciation.engine import (
        SCENARIOS,
        AssetInput,
        CalculationInput,
        compute,
    )

    report = build(SAMPLE, tax_year=2025, state=code, rules=get_rules())
    assert report["state"]["status"] == "computed"  # all three are verified for 2025
    request = CalculationInput(tax_year=2025, assets=tuple(
        AssetInput(
            id=f"r{i}", description=a.description, asset_class=a.asset_class, cost=a.cost,
            placed_in_service=a.acquired, acquired=a.acquired, recovery_years=a.recovery_years,
        ) for i, a in enumerate(SAMPLE)
    ), method="section_179_bonus", state=code)
    scenario = SCENARIOS["section_179_bonus"]
    outcome = evaluate_state(request, get_rules(), scenario, compute(replace(request, state=None), get_rules(), scenario))
    assert outcome is not None
    assert D(report["totals"]["state"]["total"]) == outcome.deduction_in(2025)


def test_federal_figures_never_change_with_the_state() -> None:
    federal_only = build(SAMPLE, tax_year=2025, state=None, rules=get_rules())
    for code in ("CA", "MN", "IL"):
        with_state = build(SAMPLE, tax_year=2025, state=code, rules=get_rules())
        assert [r["federal"] for r in with_state["rows"]] == [r["federal"] for r in federal_only["rows"]]


@pytest.mark.parametrize("code", ["IL", "TX"])
def test_no_verified_state_rule_means_no_state_figure(code: str) -> None:
    report = build(SAMPLE, tax_year=2025, state=code, rules=get_rules())
    assert report["state"]["status"] != "computed"
    assert all(r["state"] is None for r in report["rows"])
    assert report["totals"]["state"] is None


def test_an_asset_from_another_year_is_refused() -> None:
    with pytest.raises(InputError) as caught:
        build([*SAMPLE, asset("9", "Old desk", date(2024, 5, 1), "500", "furniture")],
              tax_year=2025, state=None, rules=get_rules())
    assert "not a 2025 addition" in caught.value.errors[0]


def test_more_than_the_form_limit_is_fine_for_a_report() -> None:
    many = [asset(str(i), f"Tablet {i}", date(2025, 3, 1), "1000", "computers") for i in range(150)]
    assert len(build(many, tax_year=2025, state=None, rules=get_rules())["rows"]) == 150


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------


def test_the_workbook_keeps_the_templates_columns_and_adds_the_federal_group() -> None:
    report = build(SAMPLE, tax_year=2025, state=None, rules=get_rules())
    wb = load_workbook(io.BytesIO(to_xlsx(report)))
    ws = wb.active
    assert wb.sheetnames == ["CY Addition Dep Summary"]
    assert [ws[f"{c}1"].value for c in "ABCDEFGHIJ"] == [
        "Asset #", "Asset Desription", "Date Acquired", "Cost", "Life", "Elections",
        "Federal Depreciation - Sec 179", "Federal Depreciation - Bonus",
        "Federal Depreciation - MACRS", "Federal Depreciation - Total",
    ]
    assert ws["K1"].value is None  # no state chosen: no state group
    assert ws["A2"].value == 101 and ws["C2"].value == datetime(2025, 4, 30)
    # 101 costs 12,000.50: the engine's whole-dollar amounts, not formulas.
    assert ws["G2"].value == 12001 and ws["H2"].value == 0
    # 102: 179 takes the whole cost; the formulas show the order.
    assert ws["F3"].value == "179" and ws["G3"].value == "=+D3"
    assert ws["H3"].value == "=(D3-G3)*40%"
    assert ws["I3"].value == "=(D3-G3-H3)*20%"
    assert ws["J3"].value == "=SUM(G3:I3)"
    # 105 (39-year): MACRS only. 106 (land improvements): bonus on the cost, no 179.
    assert ws["F6"].value == "MACRS only" and ws["I6"].value == "=(D6-G6-H6)*1.391%"
    assert ws["G7"].value == 0 and ws["H7"].value == "=(D7-G7)*100%"
    # Totals five rows below the last asset, as in the template.
    assert ws["D12"].value == "=SUM(D2:D7)" and ws["J12"].value == "=SUM(J2:J7)"
    assert ws["D12"].border.bottom.style == "double"
    # A note says what was applied.
    assert "Section 179, then bonus" in ws["B14"].value
    # The "Federal" block: cost and first-year depreciation by life.
    assert ws["L3"].value == "Federal" and "L3:N3" in ws.merged_cells
    assert [ws[f"L{r}"].value for r in range(5, 9)] == [5, 7, 15, 39]
    assert ws["M5"].value == "=SUMIF(E2:E7,5,D2:D7)"
    assert ws["N8"].value == "=SUMIF(E2:E7,39,J2:J7)"
    assert ws["N9"].value == "=SUM(N5:N8)"
    # Formats, fills and the IRS table.
    assert ws["D2"].number_format == '"$"#,##0.00;[Red]\\-"$"#,##0.00'
    assert ws["C2"].number_format == "mm-dd-yy"
    assert ws["J1"].fill.fgColor.rgb == "FF92D050"
    assert ws["L4"].fill.fgColor.theme == 7
    assert ws["B2"].font.name == "Aptos Narrow"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:J7"
    assert len(ws._images) == 1
    assert wb.calculation.fullCalcOnLoad


def test_a_chosen_state_adds_its_group_with_the_engines_figures() -> None:
    report = build([MACHINE], tax_year=2026, state="CA", rules=get_rules())
    ws = load_workbook(io.BytesIO(to_xlsx(report))).active
    assert [ws[f"{c}1"].value for c in "KLMN"] == [
        "State Depreciation - Sec 179", "State Depreciation - Bonus",
        "State Depreciation - MACRS", "State Depreciation - Total",
    ]
    assert [ws[f"{c}2"].value for c in "KLMN"] == [25000, 0, 10718, 35718]
    assert ws["N7"].value == "=SUM(N2:N2)"
    assert ws["P3"].value == "Federal"  # the summary block moves right of the state group
    assert any("does not allow federal bonus" in str(ws[f"B{r}"].value) for r in range(9, 13))


def test_a_state_without_a_verified_rule_leaves_its_group_empty() -> None:
    ws = load_workbook(io.BytesIO(to_xlsx(build(SAMPLE, tax_year=2025, state="IL", rules=get_rules())))).active
    assert ws["K1"].value == "State Depreciation - Sec 179"
    assert all(ws[f"{c}{r}"].value is None for c in "KLMN" for r in range(2, 8))
    assert ws["N12"].value is None
    assert any("not yet verified" in str(ws[f"B{r}"].value) for r in range(14, 18))


def test_an_election_out_is_noted_in_the_workbook() -> None:
    report = build([replace(MACHINE, elect_bonus=False, elect_179=False)], tax_year=2026, state=None,
                   rules=get_rules())
    ws = load_workbook(io.BytesIO(to_xlsx(report))).active
    assert ws["F2"].value == "MACRS only" and ws["H2"].value == 0
    assert any("Elected out of bonus for all 7-year property" in str(ws[f"B{r}"].value) for r in range(8, 12))


# ---------------------------------------------------------------------------
# Reading the upload
# ---------------------------------------------------------------------------


def test_the_templates_headers_are_read_and_its_totals_row_skipped() -> None:
    content = workbook_bytes([
        ["Asset #", "Asset Desription", "Date Acquired", "Cost", "Life"],
        [440, "Shop Equipment - Lift", datetime(2025, 4, 30), 11203.78, None],
        [441, "Copier", datetime(2025, 1, 9), "$3,438.46", 5],
        [None, None, None, None, None],
        [None, None, None, "=SUM(D2:D3)", None],
    ])
    rows = addition_import.read(content, "list.xlsx")
    assert [r.asset_id for r in rows] == ["440", "441"]
    assert rows[0].description == "Shop Equipment - Lift" and rows[0].acquired == date(2025, 4, 30)
    assert rows[1].cost == D("3438.46") and rows[1].life == "5" and rows[0].life is None


def test_columns_are_found_by_name_not_position() -> None:
    content = workbook_bytes([
        ["Notes", "Placed in service", "Description", "Amount", "Asset ID"],
        ["x", "06/23/2025", "Interior remodel", "1,000", "A-7"],
    ])
    (row,) = addition_import.read(content, "list.xlsx")
    assert (row.asset_id, row.description, row.acquired, row.cost) == ("A-7", "Interior remodel", date(2025, 6, 23), D("1000.00"))


def test_csv_and_excel_serial_dates_are_read() -> None:
    content = b"Asset #,Description,Date Acquired,Cost\n1,Tablet,45747,500\n"
    (row,) = addition_import.read(content, "list.csv")
    assert row.acquired == date(2025, 3, 31)


def test_unusable_rows_say_why_instead_of_being_guessed() -> None:
    content = workbook_bytes([
        ["Asset #", "Description", "Date Acquired", "Cost"],
        [1, "Desk", "not a date", None],
    ])
    (row,) = addition_import.read(content, "list.xlsx")
    assert row.problems == ["no readable date acquired", "no cost"]


@pytest.mark.parametrize(
    ("content", "name", "message"),
    [
        (b"x", "list.pdf", "Excel workbook"),
        (b"not a zip", "list.xlsx", "could not be opened"),
    ],
)
def test_a_file_that_is_not_an_asset_list_is_refused(content: bytes, name: str, message: str) -> None:
    with pytest.raises(addition_import.ImportError_, match=message):
        addition_import.read(content, name)


def test_a_missing_cost_column_is_named() -> None:
    with pytest.raises(addition_import.ImportError_, match="Cost"):
        addition_import.read(workbook_bytes([["Asset #", "Description", "Date Acquired"], [1, "Desk", "1/1/2025"]]), "a.xlsx")


# ---------------------------------------------------------------------------
# Life prediction: the only model call, and it only picks a class
# ---------------------------------------------------------------------------


class FakeGemini:
    def __init__(self, replies: list[Any], enabled: bool = True) -> None:
        self.replies = list(replies)
        self.enabled = enabled
        self.calls: list[dict[str, Any]] = []

    def generate(self, **kwargs: Any) -> LlmResponse:
        self.calls.append(kwargs)
        item = self.replies.pop(0)
        if isinstance(item, Exception):
            raise item
        return LlmResponse(text=json.dumps({"assets": item}), model_name="fake")


def predictor(gemini: FakeGemini) -> LifePredictor:
    return LifePredictor(gemini, get_settings(), get_rules())  # type: ignore[arg-type]


def test_the_model_picks_a_class_and_the_life_follows_from_it() -> None:
    gemini = FakeGemini([[
        {"index": 1, "asset_class": "other", "recovery_years": "5", "confidence": "high", "reason": "57.0 dealership shop equipment"},
        {"index": 2, "asset_class": "qip", "recovery_years": "", "confidence": "high", "reason": "interior remodel"},
    ]])
    predictions, message = predictor(gemini).predict(
        ["Shop Equipment - Tire Changer", "Interior remodel"], [None, None], tax_year=2026)
    assert [(p.asset_class, p.life, p.source) for p in predictions] == [("other", "5", "ai"), ("qip", "15", "ai")]
    assert message == ""
    schema = gemini.calls[0]["response_json_schema"]
    assert "qualified_production" not in schema["properties"]["assets"]["items"]["properties"]["asset_class"]["enum"]
    assert "2026 additions" in gemini.calls[0]["prompt"]


def test_lives_are_decided_for_the_chosen_year_and_jurisdiction() -> None:
    """Life prediction runs after step 1: an unsupported year or jurisdiction stops it before any model call."""
    gemini = FakeGemini([])
    with pytest.raises(UnsupportedYearError):
        predictor(gemini).predict(["Forklift"], [None], tax_year=2019)
    with pytest.raises(ValueError, match="jurisdiction"):
        predictor(gemini).predict(["Forklift"], [None], tax_year=2026, jurisdiction="ZZ")
    assert gemini.calls == []


def test_the_life_is_the_same_in_the_federal_and_state_groups() -> None:
    """No state in the rule data changes a recovery period: one life drives both groups."""
    report = build([replace(MACHINE, elect_179=False, elect_bonus=False)], tax_year=2026, state="CA",
                   rules=get_rules())
    (row,) = report["rows"]
    assert row["life"] == "7"
    assert row["federal"]["macrs"] == "14290.00" and row["state"]["macrs"] == "14290"
    assert "California follows them" in report["life_basis"]


@pytest.mark.security
def test_only_redacted_descriptions_reach_the_model() -> None:
    gemini = FakeGemini([[{"index": 1, "asset_class": "computers", "recovery_years": "", "confidence": "high", "reason": "x"}]])
    predictor(gemini).predict(["Laptop for Mr. Patel, SSN 412-88-7364"], [None], tax_year=2026)
    prompt = gemini.calls[0]["prompt"]
    assert "Patel" not in prompt and "412-88-7364" not in prompt
    assert "[REDACTED]" in prompt
    assert gemini.calls[0]["context"] == "depreciation_life"


def test_a_life_already_in_the_file_needs_no_model_call() -> None:
    gemini = FakeGemini([])
    (prediction,), _ = predictor(gemini).predict(["Land Improvement - Parking Lot"], ["15"], tax_year=2026)
    assert (prediction.asset_class, prediction.source) == ("land_improvements", "file")
    assert gemini.calls == []


def test_a_period_the_description_states_beats_the_model() -> None:
    gemini = FakeGemini([[{"index": 1, "asset_class": "qip", "recovery_years": "", "confidence": "high", "reason": "x"}]])
    (prediction,), _ = predictor(gemini).predict(["Building Improvement - Remodel (39-Year)"], [None], tax_year=2026)
    assert (prediction.asset_class, prediction.life, prediction.source) == ("nonresidential", "39", "description")


@pytest.mark.parametrize("failure", [LlmError("busy", details={"retryable": True}), ValueError("bad json")])
def test_a_failed_prediction_leaves_the_life_to_the_user(failure: Exception) -> None:
    (prediction,), message = predictor(FakeGemini([failure])).predict(["Forklift"], [None], tax_year=2026)
    assert prediction.asset_class is None and prediction.source == "none"
    assert "choose" in message


def test_an_unknown_class_from_the_model_is_not_used() -> None:
    gemini = FakeGemini([[{"index": 1, "asset_class": "spaceship", "recovery_years": "", "confidence": "high", "reason": "x"}]])
    (prediction,), _ = predictor(gemini).predict(["Thing"], [None], tax_year=2026)
    assert prediction.asset_class is None


def test_without_a_model_the_user_is_told_to_choose() -> None:
    (prediction,), message = predictor(FakeGemini([], enabled=False)).predict(["Forklift"], [None], tax_year=2026)
    assert prediction.source == "none" and "not available" in message


@pytest.mark.parametrize(
    ("life", "description", "expected"),
    [
        ("39", "Building improvement", ("nonresidential", None)),
        ("27.5", "Rental building", ("residential_rental", None)),
        ("15", "Land Improvement - Paving", ("land_improvements", None)),
        ("15", "Interior remodel", ("qip", None)),
        ("7", "Furniture - Desk", ("furniture", None)),
        ("7", "Hoist", ("machinery", None)),
        ("5", "Vehicle - Truck (GVWR 9,000)", ("heavy_vehicle", None)),
        ("5", "Vehicle - Sedan GVWR 4,500 lb", ("passenger_vehicle", None)),
        ("5", "Technology Equipment - Tablet", ("computers", None)),
        ("5", "Shop Equipment - Tire Changer", ("other", "5")),
        ("20", "Farm building", ("other", "20")),
        ("12", "Odd", (None, None)),
    ],
)
def test_a_given_life_maps_to_a_class_without_a_model(life: str, description: str, expected: tuple[Any, Any]) -> None:
    assert class_for_life(life, description) == expected


# ---------------------------------------------------------------------------
# The HTTP API
# ---------------------------------------------------------------------------


def _upload(client: TestClient, content: bytes, **data: str) -> Any:
    return client.post(
        "/api/depreciation/additions/import",
        files={"file": ("list.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data=data,
    )


LIST = workbook_bytes([
    ["Asset #", "Asset Desription", "Date Acquired", "Cost", "Life"],
    [1, "Shop Equipment - Lift", datetime(2025, 4, 30), 12000, 5],
    [2, "Land Improvement - Parking Lot", datetime(2025, 6, 23), 50000, 15],
    [3, "Mystery item", datetime(2025, 6, 23), 1000, None],
])


def test_import_with_lives_in_the_file_computes_the_report(client: TestClient) -> None:
    response = _upload(client, LIST, tax_year="2025", state="WI")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["tax_year"] == 2025 and body["state"] == "WI"
    assert "not available" in body["message"]  # Gemini is off in tests: row 3 is the user's call
    assert [r["prediction"]["source"] for r in body["rows"]] == ["file", "file", "none"]
    assert [r["asset_id"] for r in body["report"]["rows"]] == ["1", "2"]


def test_import_needs_the_tax_year_first(client: TestClient) -> None:
    """Step 1 comes first: no year means no upload, never a year guessed from the file."""
    response = _upload(client, LIST, state="WI")
    assert response.status_code == 422
    assert "tax_year" in response.text


def test_import_refuses_assets_from_another_year(client: TestClient) -> None:
    body = _upload(client, LIST, tax_year="2026").json()
    assert all("not a 2026 addition" in " ".join(r["problems"]) for r in body["rows"])
    assert body["report"] is None


def test_lives_can_be_decided_again_without_a_reupload(client: TestClient) -> None:
    payload = {"tax_year": 2026, "state": "CA", "assets": [
        {"description": "Land Improvement - Paving", "life_in_file": "15"},
        {"description": "Forklift", "life_in_file": None},
    ]}
    body = client.post("/api/depreciation/additions/lives", json=payload).json()
    assert [p["source"] for p in body["predictions"]] == ["file", "none"]
    assert body["tax_year"] == 2026 and body["state"] == "CA"
    bad = client.post("/api/depreciation/additions/lives", json={**payload, "tax_year": 2019})
    assert bad.status_code == 422


def test_report_and_export_recompute_from_the_inputs(client: TestClient) -> None:
    payload = {
        "tax_year": 2025, "state": "WI",
        "assets": [{"asset_id": "1", "description": "Copier", "acquired": "2025-01-09", "cost": "3000",
                    "asset_class": "computers", "elect_179": False}],
    }
    report = client.post("/api/depreciation/additions/report", json=payload).json()
    assert report["rows"][0]["federal"]["bonus"] == "1200.00"
    exported = client.post("/api/depreciation/additions/export", json=payload)
    assert exported.status_code == 200
    assert exported.headers["content-disposition"].endswith('CY-Addition-Dep-Summary-2025.xlsx"')
    assert load_workbook(io.BytesIO(exported.content)).active["H2"].value == "=(D2-G2)*40%"

    no_bonus = {**payload, "assets": [{**payload["assets"][0], "elect_bonus": False}]}
    report = client.post("/api/depreciation/additions/report", json=no_bonus).json()
    assert report["rows"][0]["federal"]["bonus"] == "0.00"
    assert report["rows"][0]["federal"]["macrs"] == "600.00"


def test_report_errors_come_back_as_a_422(client: TestClient) -> None:
    payload = {"tax_year": 2025, "state": "ZZ", "assets": [
        {"asset_id": "1", "description": "x", "acquired": "2025-01-09", "cost": "3000", "asset_class": "computers"}]}
    assert client.post("/api/depreciation/additions/report", json=payload).status_code == 422


def test_report_module_exposes_the_template_headers() -> None:
    assert addition_report.HEADERS["B"] == "Asset Desription"


def test_a_problem_names_its_column() -> None:
    content = workbook_bytes([
        ["Asset #", "Description", "Date Acquired", "Cost", "Life"],
        [1, "Desk", "not a date", -5, "seven"],
    ])
    (row,) = addition_import.read(content, "list.xlsx")
    assert [(i.column, i.message) for i in row.issues] == [
        ("Date Acquired", "no readable date acquired"),
        ("Cost", "cost must be more than zero"),
        ("Life", "life 'seven' is not a number"),
    ]


def test_the_template_round_trips_through_the_reader(client: TestClient) -> None:
    response = client.get("/api/depreciation/additions/template")
    assert response.status_code == 200
    ws = load_workbook(io.BytesIO(response.content)).active
    assert [ws.cell(1, c).value for c in range(1, 6)] == list(addition_import.TEMPLATE_HEADERS)
    (row,) = addition_import.read(response.content, "template.xlsx")  # the example row reads, and says what is missing
    assert {i.column for i in row.issues} == {"Date Acquired", "Cost"}


def test_import_rows_carry_file_issues_by_column(client: TestClient) -> None:
    content = workbook_bytes([
        ["Asset #", "Asset Desription", "Date Acquired", "Cost"],
        [1, "Desk", "soon", 500],
    ])
    body = _upload(client, content, tax_year="2026").json()
    assert body["rows"][0]["issues"] == [{"column": "Date Acquired", "message": "no readable date acquired"}]


def test_the_income_limit_warning_is_not_repeated_on_every_import() -> None:
    """An asset list has no business income; the screen states that once, not as a warning."""
    from app.services.depreciation.engine import NO_INCOME_LIMIT_WARNING

    report = build([MACHINE], tax_year=2026, state=None, rules=get_rules())
    assert NO_INCOME_LIMIT_WARNING not in report["warnings"]
