"""The current-year additions summary: one row per asset, in the firm's layout.

This follows the "CY Addition Dep Summary" workpaper the firm already uses
(an Excel template supplied by the owner, 2026-10-09). Its first five columns
are kept exactly - downstream workpapers match on them, "Desription" spelling
included:

    Asset # | Asset Desription | Date Acquired | Cost | Life

The template's depreciation columns were three *alternative* views of the
full cost (bonus on all of it, MACRS on all of it, state Sec. 179 on all of
it). On the owner's decision of 2026-10-11 they are replaced by the preparer's
per-asset selections, applied in this order, per jurisdiction:

    Elections | Federal: Sec 179 | Bonus | MACRS | Total
              | State:   Sec 179 | Bonus | MACRS | Total   (when a state is chosen)

1. **Section 179** on the assets the preparer elected it for, within the
   year's dollar limit and phase-out (shared by every asset in the year).
2. **Bonus** (Sec. 168(k)) on the basis left after Section 179, for the
   assets it is elected for and eligible for, at the rate the acquisition
    date and placed-in-service year give. Mixed selections within a recovery
    class are planning figures, flagged for review before filing because an
    election out on a return is class-wide (IRC 168(k)(7)).
3. **MACRS** first-year depreciation on whatever basis is left.

The elections are made on the federal return. The state group applies the
selected state's verified regime to the same elections - its own Section 179
limit, no bonus or a partial bonus, an add-back of federal bonus - through
the same engine (``state.evaluate_state``). A state without a verified rule
for the year gets no figures at all (CLAUDE.md rule 9), never federal numbers
relabelled.

Every figure comes from the deterministic engine - never from a model. The
life (MACRS class) of each asset is decided before this runs, for the tax
year and jurisdiction chosen first (``life_predictor``); in this
calculator's rule data no state changes a recovery period, so the federal
class applies in both groups and the state differs only in 179 and bonus.

The workbook carries the arithmetic as formulas where a plain rate expresses
it (``=+D2``, ``=(D2-G2)*100%``, ``=(D2-G2-H2)*14.29%``, ``SUMIF`` by life),
so a reviewer sees how each cell was produced; where the engine applied a
limit a rate cannot express (a Section 179 allocation, a Sec. 280F cap, a
state add-back) the cell holds the engine's amount instead.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, replace
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

from app.services.depreciation.engine import (
    ELECT_OUT_WARNING,
    HUNDRED,
    NO_INCOME_LIMIT_WARNING,
    SCENARIOS,
    ZERO,
    AssetInput,
    AssetResult,
    CalculationInput,
    InputError,
    compute,
    dollars,
    validate,
)
from app.services.depreciation.rules import DepreciationRules
from app.services.depreciation.state import StateOutcome, evaluate_state

CENT = Decimal("0.01")
_ASSETS = Path(__file__).parent / "report_assets"

#: The template's sheet name and its first five headers, exactly as supplied
#: - including the "Desription" spelling, because downstream workpapers
#: match on it.
SHEET_TITLE = "CY Addition Dep Summary"
HEADERS = {
    "A": "Asset #",
    "B": "Asset Desription",
    "C": "Date Acquired",
    "D": "Cost",
    "E": "Life",
    "F": "Elections",
    "G": "Federal Depreciation - Sec 179",
    "H": "Federal Depreciation - Bonus",
    "I": "Federal Depreciation - MACRS",
    "J": "Federal Depreciation - Total",
}
STATE_HEADERS = {
    "K": "State Depreciation - Sec 179",
    "L": "State Depreciation - Bonus",
    "M": "State Depreciation - MACRS",
    "N": "State Depreciation - Total",
}
#: How far below the last asset the totals row sits in the template.
TOTALS_GAP = 5
#: Most assets one report takes (all of them share one Section 179 limit and
#: one mid-quarter test, so they cannot be split across requests).
MAX_ASSETS = 500
#: Both elections, in the engine's scenario terms; what is not elected is
#: turned off per asset.
_SCENARIO = SCENARIOS["section_179_bonus"]


@dataclass(frozen=True, slots=True)
class AdditionAsset:
    """One asset of the report, with its class (life) already decided."""

    asset_id: str
    description: str
    acquired: date
    cost: Decimal
    asset_class: str
    #: Only for the "other" class, whose recovery period is chosen.
    recovery_years: str | None = None
    #: Section 179 is elected asset by asset; unticked, the asset takes none.
    elect_179: bool = True
    #: Per-asset report selection; mixed classes require review before filing.
    elect_bonus: bool = True
    #: A heavy vehicle outside the SUV cap (cargo van, work truck, ...).
    suv_cap_exempt: bool = False


def build(
    assets: list[AdditionAsset],
    *,
    tax_year: int,
    state: str | None,
    rules: DepreciationRules,
) -> dict[str, Any]:
    """Compute every row of the report. Raises ``InputError`` / ``UnsupportedYearError``."""
    if not assets:
        raise InputError(["Add at least one asset."])
    if len(assets) > MAX_ASSETS:
        raise InputError([f"At most {MAX_ASSETS} assets per report."])

    engine_assets = tuple(
        AssetInput(
            id=f"r{index}",
            description=asset.description,
            asset_class=asset.asset_class,
            cost=asset.cost,
            placed_in_service=asset.acquired,
            acquired=asset.acquired,
            recovery_years=asset.recovery_years,
            section_179_amount=None if asset.elect_179 else ZERO,
            elect_out_bonus=not asset.elect_bonus,
            suv_cap_exempt=asset.suv_cap_exempt,
        )
        for index, asset in enumerate(assets)
    )
    request = CalculationInput(tax_year=tax_year, assets=engine_assets, method=_SCENARIO.key, bonus_by_asset=True)
    validate(request, rules, max_assets=MAX_ASSETS)
    out_of_year = [
        f"{a.description or a.asset_id}: acquired in {a.acquired.year}, not a {tax_year} addition."
        for a in assets if a.acquired.year != tax_year
    ]
    if out_of_year:
        raise InputError(out_of_year)

    federal = compute(request, rules, _SCENARIO)
    fed = {a.id: a for a in federal.assets}

    outcome: StateOutcome | None = None
    if state:
        outcome = evaluate_state(replace(request, state=state), rules, _SCENARIO, federal)
    st = (
        {a.id: a for a in outcome.computation.assets}
        if outcome is not None and outcome.status == "computed" and outcome.computation is not None
        else None
    )
    addbacks = _bonus_addbacks(outcome, tax_year) if st is not None else {}

    class_choices: dict[str, set[bool]] = {}
    for index, asset in enumerate(assets):
        if rules.classes[asset.asset_class].bonus_eligible:
            class_choices.setdefault(fed[f"r{index}"].recovery_years, set()).add(asset.elect_bonus)
    out_classes = sorted(
        (period for period, selections in class_choices.items() if selections == {False}),
        key=Decimal,
    )
    election_warnings = [
        (f"Planning only: mixed bonus selections for {period}-year property in {tax_year}. "
         "Figures follow each asset's selection. Before filing, reconcile the class-wide "
         "bonus election required by IRC 168(k)(7).")
        for period, selections in sorted(class_choices.items(), key=lambda item: Decimal(item[0]))
        if len(selections) > 1
    ]

    rows: list[dict[str, Any]] = []
    #: Unrounded federal amounts, as the workbook's formulas produce them.
    #: Totals and the by-life summary add these and round once - as Excel's
    #: SUM does over the formula cells - so screen and file agree to the cent.
    exact: list[dict[str, Decimal]] = []
    for index, asset in enumerate(assets):
        key = f"r{index}"
        f = fed[key]
        cls = rules.classes[asset.asset_class]
        figures = _federal_figures(f, asset.cost, tax_year)
        exact.append(figures)
        bonus_on = asset.elect_bonus
        row: dict[str, Any] = {
            "asset_id": asset.asset_id,
            "description": asset.description,
            "acquired": asset.acquired.isoformat(),
            "cost": str(asset.cost),
            "life": f.recovery_years,
            "asset_class": f.asset_class,
            "class_label": f.class_label,
            "convention": f.convention,
            "elect_179": asset.elect_179,
            "elect_bonus": asset.elect_bonus,
            "sec179_eligible": cls.section_179_eligible,
            "sec179_reason": _sec179_reason(cls),
            "bonus_eligible": cls.bonus_eligible,
            "bonus_reason": _bonus_reason(cls),
            "bonus_applied": bool(cls.bonus_eligible and bonus_on),
            "bonus_rate": _pct(f.bonus_rate),
            "elections": _elections(f, cls.bonus_eligible and bonus_on),
            "federal": {
                "sec179": str(_cents(figures["sec179"])),
                "bonus": str(_cents(figures["bonus"])),
                "macrs": str(_cents(figures["macrs"])),
                "total": str(_cents(figures["total"])),
                "sec179_full": figures["sec179_full"] == 1,
                "bonus_formula": figures["bonus_formula"] == 1,
                "macrs_formula": figures["macrs_formula"] == 1,
                "macrs_rate": _pct(f.rows[0].rate or ZERO),
            },
            "state": None,
        }
        if st is not None:
            s = st[key]
            first = s.row_for(tax_year) or s.rows[0]
            added = addbacks.get(key, ZERO)
            row["state"] = {
                "sec179": str(first.section_179),
                "bonus": str(first.bonus - added),
                "macrs": str(first.macrs),
                "total": str(first.total - added),
                "addback": str(added),
            }
        rows.append(row)

    totals: dict[str, Any] = {
        "cost": str(sum((Decimal(r["cost"]) for r in rows), ZERO)),
        "federal": {k: str(_cents(sum((e[k] for e in exact), ZERO))) for k in ("sec179", "bonus", "macrs", "total")},
        "state": None,
    }
    if st is not None and outcome is not None:
        state_totals = {
            k: sum((Decimal(r["state"][k]) for r in rows), ZERO) for k in ("sec179", "bonus", "macrs", "total")
        }
        # A Section 179 add-back is made on the year's total, not per asset.
        adjustment = outcome.deduction_in(tax_year) - state_totals["total"]
        state_totals["adjustment"] = adjustment
        state_totals["total"] += adjustment
        totals["state"] = {k: str(v) for k, v in state_totals.items()}

    return {
        "tax_year": tax_year,
        "state": _state_info(outcome, rules, tax_year),
        "rows": rows,
        "totals": totals,
        "summary": _summary(rows, [e["total"] for e in exact]),
        "elect_out_classes": out_classes,
        "section_179": _section_179_summary(federal, tax_year),
        # An asset list carries no business income: the income-limit warning
        # would fire on every import, so the screen states it once instead.
        # The election out is reported as elect_out_classes, not repeated here.
        "warnings": list(dict.fromkeys(
            w for w in [*election_warnings, *federal.warnings]
            if w != NO_INCOME_LIMIT_WARNING and not w.startswith(ELECT_OUT_WARNING)
        )),
        "life_basis": (
            f"Lives are federal MACRS recovery periods (Rev. Proc. 87-56) for {tax_year}. "
            + (
                f"{outcome.name} follows them: in this calculator's rule data no state changes a "
                "recovery period, only Section 179 and bonus."
                if outcome is not None else ""
            )
        ).strip(),
    }


def _federal_figures(f: AssetResult, cost: Decimal, tax_year: int) -> dict[str, Decimal]:
    """First-year Sec 179, bonus and MACRS as the workbook's cells produce them.

    The template's formulas work on the cost column (``=(D2-G2)*40%``). The
    engine works on the basis rounded to whole dollars, as a return does, so
    the formulas are used only when the two are the same number - a cost in
    whole dollars - and then the screen shows the formula's unrounded result
    (totals round once, as Excel's SUM does). Otherwise, and wherever a limit
    a rate cannot express applies (a Section 179 allocation, a 280F cap), the
    cell and the screen both hold the engine's amount. ``*_formula`` is 1
    when the workbook may write the formula.
    """
    first = f.row_for(tax_year) or f.rows[0]
    sec179 = first.section_179
    whole = f.business_basis == cost
    after_179 = f.business_basis - f.section_179_elected
    bonus_formula = whole and f.bonus == dollars(after_179 * f.bonus_rate / HUNDRED)
    bonus = after_179 * f.bonus_rate / HUNDRED if bonus_formula else f.bonus
    macrs_formula = (whole and bonus_formula and not first.capped and not first.special
                     and first.base == f.macrs_basis and first.rate is not None)
    # The MACRS cell is (cost - Sec 179 cell - bonus cell) x rate, the bonus
    # cell unrounded; the screen uses the same basis so the two agree to the
    # cent (the engine's own schedule rounds bonus to dollars first).
    macrs = (cost - sec179 - bonus) * (first.rate or ZERO) / HUNDRED if macrs_formula else first.macrs
    return {
        "sec179": sec179,
        "bonus": bonus,
        "macrs": macrs,
        "total": sec179 + bonus + first.special + macrs,
        "sec179_full": Decimal(1 if whole and sec179 == cost else 0),
        "bonus_formula": Decimal(1 if bonus_formula else 0),
        "macrs_formula": Decimal(1 if macrs_formula else 0),
    }


def _bonus_addbacks(outcome: StateOutcome | None, tax_year: int) -> dict[str, Decimal]:
    """Per asset: the part of federal bonus the state adds back this year."""
    if outcome is None or outcome.computation is None or outcome.entry is None:
        return {}
    rule = outcome.entry.get("bonus") or {}
    if rule.get("mode") != "addback":
        return {}
    pct = Decimal(str(rule["addback_pct"]))
    return {
        a.id: dollars(a.bonus * pct / HUNDRED)
        for a in outcome.computation.assets
        if a.placed_in_service.year == tax_year and a.bonus > 0
    }


def _state_info(outcome: StateOutcome | None, rules: DepreciationRules, tax_year: int) -> dict[str, Any] | None:
    if outcome is None:
        return None
    entry = outcome.entry or {}
    bonus = entry.get("bonus") or {}
    mode = bonus.get("mode", "")
    name = outcome.name
    if mode == "none":
        bonus_note = (f"{name} does not allow federal bonus depreciation for {tax_year}; "
                      "the basis is depreciated under MACRS instead.")
    elif mode == "percent":
        after = bonus.get("applies_if_acquired_after")
        bonus_note = (f"{name} allows {bonus.get('allowed_pct')}% bonus"
                      + (f" for property acquired after {after}" if after else "") + f" in {tax_year}.")
    elif mode == "pre_obbba":
        bonus_note = f"{name} follows the bonus rates in effect before the 2025 federal law."
    elif mode == "addback":
        bonus_note = (f"{name} adds back {bonus.get('addback_pct')}% of federal bonus and subtracts it "
                      f"over the next {bonus.get('recovery_years')} years.")
    elif mode == "follow":
        bonus_note = f"{name} follows federal bonus depreciation for {tax_year}."
    else:
        bonus_note = ""
    return {
        "code": outcome.code,
        "name": name,
        "status": outcome.status,
        "message": outcome.message,
        "bonus_mode": mode,
        "bonus_allowed": mode not in ("none", "not_applicable", "unknown", ""),
        "bonus_note": bonus_note,
        "section_179_mode": (entry.get("section_179") or {}).get("mode", ""),
        "source_url": entry.get("source_url", ""),
        "source_citation": entry.get("source_citation", ""),
        "notes": outcome.notes,
    }


def _section_179_summary(federal: Any, tax_year: int) -> dict[str, str] | None:
    summary = federal.section_179.get(tax_year)
    if summary is None:
        return None
    return {
        "limit": str(summary.limit),
        "threshold": str(summary.threshold) if summary.threshold is not None else "",
        "total_cost": str(summary.total_cost),
        "reduction": str(summary.reduction),
        "dollar_limit": str(summary.dollar_limit),
        "elected": str(summary.elected),
        "deduction": str(summary.deduction),
    }


def _sec179_reason(cls: Any) -> str:
    if cls.section_179_eligible:
        return ""
    if cls.key == "land_improvements":
        return "Land improvements are not section 179 property."
    if cls.real_property:
        return "Buildings are not section 179 property."
    return "Not section 179 property."


def _bonus_reason(cls: Any) -> str:
    if cls.bonus_eligible:
        return ""
    if cls.qpp:
        return "Qualified production property uses the section 168(n) allowance, not bonus."
    if cls.real_property:
        return f"{cls.recovery_years}-year property: the recovery period is over 20 years, so it is not qualified property."
    return "Not qualified property for bonus depreciation."


def _elections(f: AssetResult, bonus_on: bool) -> str:
    parts = []
    if f.section_179_allowed > 0:
        parts.append("179")
    if bonus_on and f.bonus > 0:
        parts.append("Bonus")
    return "+".join(parts) or "MACRS only"


def _summary(rows: list[dict[str, Any]], totals_exact: list[Decimal]) -> list[dict[str, str]]:
    """Cost and federal first-year depreciation by life - the template's "Federal" block."""
    lives = sorted({r["life"] for r in rows}, key=Decimal)
    return [
        {
            "life": life,
            "cost": str(sum((Decimal(r["cost"]) for r in rows if r["life"] == life), ZERO)),
            "depreciation": str(_cents(sum(
                (amount for r, amount in zip(rows, totals_exact) if r["life"] == life), ZERO
            ))),
        }
        for life in lives
    ]


def _cents(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------


def to_xlsx(report: dict[str, Any]) -> bytes:
    """The report as the template workbook: same sheet, A-E columns, formats and formulas."""
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.drawing.image import Image  # type: ignore[import-untyped]
    from openpyxl.styles import (  # type: ignore[import-untyped]
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.styles.colors import Color  # type: ignore[import-untyped]
    from openpyxl.workbook import properties  # type: ignore[import-untyped]

    wb = Workbook()
    # The template's own Office theme, so its theme colours (the summary
    # block's fill) render exactly as they do there.
    wb.loaded_theme = (_ASSETS / "office_theme.xml").read_bytes()
    normal = next(style for style in wb._named_styles if style.name == "Normal")
    normal.font = Font(name="Aptos Narrow", size=11)
    # Formulas carry no cached values when written here: compute them on open.
    wb.calculation = properties.CalcProperties(fullCalcOnLoad=True)

    ws = wb.active
    ws.title = SHEET_TITLE
    ws.sheet_properties.tabColor = Color(theme=6)

    font = Font(name="Aptos Narrow", size=11)
    bold = Font(name="Aptos Narrow", size=11, bold=True)
    italic = Font(name="Aptos Narrow", size=10, italic=True)
    thin = Side(style="thin")
    green = PatternFill(fill_type="solid", fgColor="FF92D050")
    summary_fill = PatternFill(fill_type="solid", fgColor=Color(theme=7, tint=0.7999816888943144))
    money = '"$"#,##0.00;[Red]\\-"$"#,##0.00'
    accounting = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'

    state: dict[str, Any] | None = report.get("state")
    with_state = state is not None
    headers = {**HEADERS, **(STATE_HEADERS if with_state else {})}
    figure_columns = "GHIJ" + ("KLMN" if with_state else "")
    last_column = "N" if with_state else "J"

    for column, header in headers.items():
        cell = ws[f"{column}1"]
        cell.value = header
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws["J1"].fill = green

    rows = report["rows"]
    first, last = 2, 1 + len(rows)
    for offset, row in enumerate(rows):
        r = first + offset
        fed = row["federal"]
        ws[f"A{r}"] = _id_value(row["asset_id"])
        ws[f"B{r}"] = row["description"]
        ws[f"B{r}"].alignment = Alignment(wrap_text=True)
        ws[f"C{r}"] = datetime.fromisoformat(row["acquired"])
        ws[f"C{r}"].number_format = "mm-dd-yy"
        ws[f"D{r}"] = float(Decimal(row["cost"]))
        ws[f"D{r}"].number_format = money
        ws[f"E{r}"] = _life_value(row["life"])
        ws[f"F{r}"] = row["elections"]
        ws[f"G{r}"] = f"=+D{r}" if fed["sec179_full"] else float(Decimal(fed["sec179"]))
        ws[f"H{r}"] = _bonus_cell(row, r)
        ws[f"I{r}"] = _macrs_cell(row, r)
        ws[f"J{r}"] = f"=SUM(G{r}:I{r})"
        if with_state:
            s = row.get("state")
            for column, key in zip("KLMN", ("sec179", "bonus", "macrs", "total")):
                # No verified state rule: no figure, never another regime's.
                ws[f"{column}{r}"] = float(Decimal(s[key])) if s else None
        for column in "ABCDEF" + figure_columns:
            ws[f"{column}{r}"].font = font
        for column in "E" + figure_columns:
            ws[f"{column}{r}"].number_format = accounting

    totals = last + TOTALS_GAP
    for column in "D" + figure_columns:
        cell = ws[f"{column}{totals}"]
        if column == "N" and with_state and report["totals"]["state"] is not None:
            adjustment = Decimal(report["totals"]["state"].get("adjustment", "0"))
            cell.value = f"=SUM(N{first}:N{last})" + (f"{adjustment:+}" if adjustment else "")
        elif column in "KLMN" and report["totals"]["state"] is None:
            cell.value = None
        else:
            cell.value = f"=SUM({column}{first}:{column}{last})"
        cell.font = bold
        cell.number_format = accounting
        cell.border = Border(top=thin, bottom=Side(style="double"))
    ws[f"J{totals}"].fill = green

    # What was applied, where: the jurisdiction, the year and the rule status.
    notes = [(f"Tax year {report['tax_year']}. Federal: Section 179, then bonus on the basis left, "
              "then MACRS on the rest, per the elections in column F.")]
    if state is not None:
        notes.append(f"{state['name']}: {state['message']}")
        if state.get("bonus_note"):
            notes.append(state["bonus_note"])
        adjustment = Decimal((report["totals"]["state"] or {}).get("adjustment", "0"))
        if adjustment:
            notes.append(f"{state['name']} Section 179 adjustment on the year's total: {adjustment:,.2f} "
                         "(included in the state total).")
    if report.get("elect_out_classes"):
        notes.append("Elected out of bonus for all "
                     + ", ".join(f"{p}-year" for p in report["elect_out_classes"])
                     + " property placed in service this year (the election is made by class).")
    notes.extend(report.get("warnings", []))
    for n, text in enumerate(notes):
        ws[f"B{totals + 2 + n}"] = text
        ws[f"B{totals + 2 + n}"].font = italic

    summary_col = "P" if with_state else "L"
    _summary_block(ws, report["summary"], first, last, bold, font, summary_fill, accounting, thin, summary_col)
    summary_end = 4 + len(report["summary"]) + 1
    image = Image(str(_ASSETS / "irs_pub946_table_a1.png"))
    image.width, image.height = 513, 411
    ws.add_image(image, f"{summary_col}{summary_end + 1}")

    widths = {"B": 59.21875, "C": 13.88671875, "D": 13.33203125, "E": 9.0, "F": 12.5}
    for column in figure_columns:
        widths[column] = 18.0
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_column}{last}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _summary_block(ws: Any, summary: list[dict[str, str]], first: int, last: int,
                   bold: Any, font: Any, fill: Any, accounting: str, thin: Any, start: str) -> None:
    from openpyxl.styles import Alignment, Border  # type: ignore[import-untyped]
    from openpyxl.utils import (  # type: ignore[import-untyped]
        column_index_from_string,
        get_column_letter,
    )

    life_col = start
    cost_col = get_column_letter(column_index_from_string(start) + 1)
    dep_col = get_column_letter(column_index_from_string(start) + 2)
    ws.merge_cells(f"{life_col}3:{dep_col}3")
    ws[f"{life_col}3"] = "Federal"
    ws[f"{life_col}3"].font = bold
    ws[f"{life_col}3"].alignment = Alignment(horizontal="center")
    for column in (life_col, cost_col, dep_col):
        ws[f"{column}3"].fill = fill
    for column, label in ((life_col, "Life"), (cost_col, "Cost"), (dep_col, "Depreciation")):
        cell = ws[f"{column}4"]
        cell.value = label
        cell.font = bold
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=thin)
        cell.fill = fill

    life_range, cost_range, dep_range = f"E{first}:E{last}", f"D{first}:D{last}", f"J{first}:J{last}"
    r = 5
    for line in summary:
        life = _life_value(line["life"])
        ws[f"{life_col}{r}"] = life
        ws[f"{cost_col}{r}"] = f"=SUMIF({life_range},{life},{cost_range})"
        ws[f"{dep_col}{r}"] = f"=SUMIF({life_range},{life},{dep_range})"
        for column in (life_col, cost_col, dep_col):
            ws[f"{column}{r}"].fill = fill
            ws[f"{column}{r}"].font = font
            ws[f"{column}{r}"].number_format = accounting
        r += 1
    ws[f"{dep_col}{r}"] = f"=SUM({dep_col}5:{dep_col}{r - 1})"
    ws[f"{dep_col}{r}"].font = font
    ws[f"{dep_col}{r}"].number_format = accounting
    ws[f"{dep_col}{r}"].border = Border(top=thin)
    ws.column_dimensions[cost_col].width = 13.33203125
    ws.column_dimensions[dep_col].width = 13.33203125


def _bonus_cell(row: dict[str, Any], r: int) -> Any:
    fed = row["federal"]
    rate = Decimal(row["bonus_rate"])
    if not fed["bonus_formula"]:
        return float(Decimal(fed["bonus"]))
    if rate == 0 or not row["bonus_applied"]:
        return 0
    return f"=(D{r}-G{r})*{_pct(rate)}%"


def _macrs_cell(row: dict[str, Any], r: int) -> Any:
    fed = row["federal"]
    if not fed["macrs_formula"]:
        return float(Decimal(fed["macrs"]))
    value = Decimal(fed["macrs_rate"])
    return f"=(D{r}-G{r}-H{r})*{_pct(value)}%" if value else 0


def _pct(rate: Decimal | str) -> str:
    text = format(Decimal(rate).normalize(), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _life_value(life: str) -> int | float:
    value = Decimal(life)
    return int(value) if value == value.to_integral_value() else float(value)


def _id_value(asset_id: str) -> int | str:
    return int(asset_id) if asset_id.isdigit() else asset_id
