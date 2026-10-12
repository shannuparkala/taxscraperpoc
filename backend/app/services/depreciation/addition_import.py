"""Reading an asset list from an uploaded spreadsheet.

The expected input is the firm's additions list before anyone has worked on
it: asset number, description, date acquired and cost - the first four
columns of the "CY Addition Dep Summary" template. A Life column, when
present, is used as given (the preparer's decision), and anything else in
the sheet is ignored.

Columns are found by their header, not their position, so "Asset #" /
"Asset ID", "Asset Desription" (the template's spelling) / "Description",
"Date Acquired" / "Placed in service", "Cost" / "Amount" all work. A row is
read only when it has a description; the template's totals row and blank
spacer rows are skipped. A row that cannot be used says why, by row number,
rather than being guessed at.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

#: Upload limits: generous for an asset list, small for anything else.
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500

_HEADERS: dict[str, tuple[str, ...]] = {
    "asset_id": ("asset#", "assetno", "assetnumber", "assetid", "asset", "id", "no", "number", "#", "tag", "assettag"),
    "description": ("assetdesription", "assetdescription", "description", "desc", "assetname", "name", "item"),
    "acquired": ("dateacquired", "acquired", "acquisitiondate", "dateplacedinservice", "placedinservice",
                 "inservicedate", "pisdate", "servicedate", "date"),
    "cost": ("cost", "costbasis", "basis", "originalcost", "amount", "price"),
    "life": ("life", "recoveryperiod", "usefullife", "macrslife", "years"),
}
_DATE_FORMATS = ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%m-%d-%y", "%d-%b-%Y", "%b %d, %Y", "%B %d, %Y")


class ImportError_(ValueError):
    """The file cannot be read as an asset list at all."""


#: The upload template's header row: the first five columns of the firm's
#: "CY Addition Dep Summary" workpaper, "Desription" spelling included.
TEMPLATE_HEADERS = ("Asset #", "Asset Desription", "Date Acquired", "Cost", "Life")


@dataclass(frozen=True, slots=True)
class Issue:
    """Why a row cannot be used, and which column of the upload it is in."""

    column: str
    message: str

    def to_dict(self) -> dict[str, str]:
        return {"column": self.column, "message": self.message}


@dataclass(slots=True)
class ImportedRow:
    row: int
    asset_id: str
    description: str
    acquired: date | None
    cost: Decimal | None
    life: str | None
    issues: list[Issue] = field(default_factory=list)

    @property
    def problems(self) -> list[str]:
        return [issue.message for issue in self.issues]

    @property
    def usable(self) -> bool:
        return not self.issues


def template_xlsx() -> bytes:
    """An empty upload template: the header row, one example row, and a Life column to fill or leave blank."""
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font  # type: ignore[import-untyped]

    wb = Workbook()
    ws = wb.active
    ws.title = "Additions"
    ws.append(list(TEMPLATE_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append([1, "Example - delete this row", None, None, None])
    ws["C2"].number_format = "mm/dd/yyyy"
    ws["D2"].number_format = '"$"#,##0.00'
    for column, width in {"A": 10, "B": 50, "C": 15, "D": 15, "E": 8}.items():
        ws.column_dimensions[column].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def read(content: bytes, filename: str) -> list[ImportedRow]:
    if len(content) > MAX_BYTES:
        raise ImportError_("That file is larger than 5 MB; an asset list should be far smaller.")
    name = filename.lower()
    if name.endswith(".csv"):
        grid = _csv_grid(content)
    elif name.endswith((".xlsx", ".xlsm")):
        grid = _xlsx_grid(content)
    else:
        raise ImportError_("Upload an Excel workbook (.xlsx) or a CSV file.")

    header_row, columns = _find_header(grid)
    missing = [label for key, label in (("description", "Description"), ("acquired", "Date Acquired"),
                                         ("cost", "Cost")) if key not in columns]
    if missing:
        raise ImportError_(
            "Could not find the column(s): " + ", ".join(missing)
            + ". The first sheet needs a header row with Asset #, Description, Date Acquired and Cost."
        )

    rows: list[ImportedRow] = []
    for offset, values in enumerate(grid[header_row + 1 :], start=header_row + 2):
        description = _text(_get(values, columns.get("description")))
        if not description:
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportError_(f"At most {MAX_ROWS} assets per file.")
        asset_id = _text(_get(values, columns.get("asset_id"))) or str(len(rows) + 1)
        item = ImportedRow(offset, _id_text(asset_id), description, None, None, None)
        item.acquired = _date(_get(values, columns.get("acquired")))
        if item.acquired is None:
            item.issues.append(Issue("Date Acquired", "no readable date acquired"))
        item.cost = _money(_get(values, columns.get("cost")))
        if item.cost is None:
            item.issues.append(Issue("Cost", "no cost"))
        elif item.cost <= 0:
            item.issues.append(Issue("Cost", "cost must be more than zero"))
        life = _get(values, columns.get("life"))
        if life not in (None, ""):
            try:
                item.life = format(Decimal(str(life).strip()).normalize(), "f")
            except InvalidOperation:
                item.issues.append(Issue("Life", f"life '{life}' is not a number"))
        rows.append(item)
    if not rows:
        raise ImportError_("The file has a header row but no assets under it.")
    return rows


# ---------------------------------------------------------------------------


def _xlsx_grid(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    try:
        # data_only: a formula cell is read as the value Excel last computed.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("That file could not be opened as an Excel workbook.") from exc
    sheet = workbook.worksheets[0]
    grid = [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 60)]
    workbook.close()
    return grid


def _csv_grid(content: bytes) -> list[list[Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    return [list(row) for row in csv.reader(io.StringIO(text))]


def _find_header(grid: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """The first of the top 20 rows that names a description column, and its columns."""
    for index, values in enumerate(grid[:20]):
        columns: dict[str, int] = {}
        keys = [re.sub(r"[^a-z0-9#]", "", str(value or "").lower()) for value in values]
        # Exact names first, so "Asset Desription" is never taken for "Asset #".
        for position, key in enumerate(keys):
            for field_name, names in _HEADERS.items():
                if key and field_name not in columns and key in names:
                    columns[field_name] = position
                    break
        # Then looser readings of headers no exact name matched.
        taken = set(columns.values())
        for position, key in enumerate(keys):
            if not key or position in taken:
                continue
            loose = _loose_field(key)
            if loose and loose not in columns:
                columns[loose] = position
                taken.add(position)
        if "description" in columns and len(columns) >= 2:
            return index, columns
    raise ImportError_("Could not find a header row naming the asset description in the first 20 rows.")


def _loose_field(key: str) -> str | None:
    if "descri" in key or "desription" in key:
        return "description"
    if "acqui" in key or "inservice" in key or key.startswith("date"):
        return "acquired"
    if "cost" in key or "basis" in key:
        return "cost"
    if key.startswith(("asset", "tag")) and ("#" in key or "no" in key or "id" in key or "num" in key):
        return "asset_id"
    if "life" in key or "recovery" in key:
        return "life"
    return None


def _get(values: list[Any], position: int | None) -> Any:
    if position is None or position >= len(values):
        return None
    return values[position]


def _text(value: Any) -> str:
    return "" if value is None else " ".join(str(value).split())


def _id_text(value: str) -> str:
    # Excel stores 440 as 440.0 in some exports.
    return value[:-2] if re.fullmatch(r"\d+\.0", value) else value


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:
        # An Excel serial date that arrived as a number.
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = _text(value)
    if text.isdigit():
        return _date(int(text))
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()  # noqa: DTZ007 - a calendar date, no time of day
        except ValueError:
            continue
    return None


def _money(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = _text(value)
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    try:
        amount = Decimal(cleaned).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -amount if negative else amount
